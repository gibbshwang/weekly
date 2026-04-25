from pathlib import Path
from openpyxl import load_workbook
from scripts.xlsx_template import generate_xlsx


def test_generate_xlsx_has_headers(tmp_path: Path):
    out = tmp_path / "_지시사항.xlsx"
    generate_xlsx(out_path=out, dept_name="기획팀", parts=["전략기획", "사업개발"])
    assert out.exists()
    wb = load_workbook(out)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert headers == ["일자", "지시내용", "담당파트", "우선순위", "마감", "비고"]


def test_generate_xlsx_has_part_dropdown(tmp_path: Path):
    out = tmp_path / "_지시사항.xlsx"
    generate_xlsx(out_path=out, dept_name="기획팀", parts=["전략기획", "사업개발"])
    wb = load_workbook(out)
    ws = wb.active
    dvs = ws.data_validations.dataValidation
    formulas = [dv.formula1 for dv in dvs]
    # At least one DV references all our parts
    assert any("전략기획" in f and "사업개발" in f for f in formulas)


def test_generate_xlsx_has_priority_dropdown(tmp_path: Path):
    out = tmp_path / "_지시사항.xlsx"
    generate_xlsx(out_path=out, dept_name="기획팀", parts=["전략기획"])
    wb = load_workbook(out)
    ws = wb.active
    dvs = ws.data_validations.dataValidation
    formulas = [dv.formula1 for dv in dvs]
    assert any("높음" in f and "보통" in f and "낮음" in f for f in formulas)


def test_generate_xlsx_creates_parent_dir(tmp_path: Path):
    out = tmp_path / "공유폴더" / "기획팀" / "_지시사항.xlsx"
    generate_xlsx(out_path=out, dept_name="기획팀", parts=["전략기획"])
    assert out.exists()
