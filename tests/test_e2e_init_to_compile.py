"""End-to-end Phase 2: init pipeline → prepare → assign → compile, with mocked LLM + SMTP.

Walks the full happy path:
  1. precheck + scaffold + write_config + keyring
  2. wreport prepare 팀 --prev-week ''  (first-week — no carry-forward)
  3. populate _지시사항.xlsx with two rows (담당파트 column-set + LLM-inferred)
  4. wreport assign 팀                  (rows route to part workbooks with W18-NNN IDs)
  5. fill in 처리결과 / 상태 in part workbooks
  6. wreport compile 팀                 (LLM produces v2 schema → dashboard → mail)
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from unittest.mock import MagicMock, patch

import keyring
import pytest
from openpyxl import Workbook, load_workbook

from scripts.precheck import precheck
from scripts.scope import TeamScopeAnswers, GroupScopeAnswer, PartScopeAnswer
from scripts.scaffold import scaffold_project
from scripts.config_setup import write_config
from scripts.xlsx_template import (
    generate_directives_xlsx,
    PART_SHEET_HEADERS,
)
from scripts.keyring_setup import save_secrets
from scripts.handoff import generate_readme


BUNDLE_ROOT = Path(__file__).resolve().parent.parent


class _FakeBackend(keyring.backend.KeyringBackend):
    priority = 1

    def __init__(self):
        self._store = {}

    def set_password(self, service, username, password):
        self._store[(service, username)] = password

    def get_password(self, service, username):
        return self._store.get((service, username))

    def delete_password(self, service, username):
        self._store.pop((service, username), None)


@pytest.fixture
def fake_keyring():
    keyring.set_keyring(_FakeBackend())
    yield


def _make_answers(storage_root: Path) -> TeamScopeAnswers:
    return TeamScopeAnswers(
        team_name="테스트팀",
        team_lead_name="팀장 김",
        team_lead_email="team_lead@example.com",
        groups=[
            GroupScopeAnswer(
                name="사업그룹",
                lead_name="그룹장 이",
                lead_email="biz_lead@example.com",
                parts=[
                    PartScopeAnswer(name="전략기획", lead_name="김파트", lead_email="p1@example.com"),
                    PartScopeAnswer(name="사업개발", lead_name="이파트", lead_email="p2@example.com"),
                ],
            ),
        ],
        storage_type="local",
        storage_root=str(storage_root),
        smtp_host="smtp.example.com",
        smtp_port=587,
        smtp_use_tls=True,
        smtp_user="op@example.com",
        smtp_password="pwd",
        ai_provider="codex",
    )


def _build_team_stub(answers: TeamScopeAnswers):
    """Mirror the Pydantic Team shape just enough for generate_directives_xlsx."""
    from dataclasses import dataclass
    from typing import List

    @dataclass
    class P:
        name: str

    @dataclass
    class G:
        name: str
        parts: List[P]

    @dataclass
    class T:
        name: str
        groups: List[G]
        def iter_parts(self):
            for g in self.groups:
                for p in g.parts:
                    yield g, p

    return T(
        name=answers.team_name,
        groups=[
            G(name=g.name, parts=[P(name=p.name) for p in g.parts])
            for g in answers.groups
        ],
    )


def test_init_to_compile_e2e_phase2(tmp_path: Path, fake_keyring):
    project_root = tmp_path / "weekly_test_project"
    storage_root = tmp_path / "weekly-test-data"
    answers = _make_answers(storage_root)

    # ----- Init pipeline -----
    precheck(target_dir=project_root)
    scaffold_project(
        target=project_root,
        team_slug="테스트팀_slug",
        team_name="테스트팀",
        bundle_root=BUNDLE_ROOT,
        install_pkg=False,
    )
    write_config(answers, target=project_root, bundle_root=BUNDLE_ROOT)
    save_secrets(answers)
    generate_readme(
        target=project_root, bundle_root=BUNDLE_ROOT,
        team_name="테스트팀", project_path=str(project_root),
        storage_root=str(storage_root),
    )
    team_data_root = storage_root / "테스트팀"
    generate_directives_xlsx(
        out_path=team_data_root / "_지시사항.xlsx",
        team=_build_team_stub(answers),
    )

    # Verify init artifacts exist with the new shape
    assert (project_root / "config.yaml").exists()
    assert (project_root / "weekly_runtime" / "cli.py").exists()
    assert (project_root / "weekly_runtime" / "_llm.py").exists()
    assert (project_root / "weekly_runtime" / "templates" / "template_dashboard_v2.html.j2").exists()
    assert (project_root / "weekly_runtime" / "prompts" / "compile_system_v2.txt").exists()
    assert (project_root / "weekly_runtime" / "prompts" / "assign_system.txt").exists()
    assert (team_data_root / "_지시사항.xlsx").exists()
    assert keyring.get_password("weekly-테스트팀", "smtp_password") == "pwd"
    assert (project_root / "README.md").exists()

    # Drop any cached weekly_runtime modules so the rendered package is freshly imported
    for k in list(sys.modules):
        if k == "weekly_runtime" or k.startswith("weekly_runtime."):
            del sys.modules[k]

    sys.path.insert(0, str(project_root))
    try:
        import weekly_runtime._llm as _llm_mod

        from click.testing import CliRunner
        from weekly_runtime.cli import cli as wreport_cli

        # ----- Step 2: wreport prepare (first week, no carry-forward) -----
        runner = CliRunner()
        prep_result = runner.invoke(
            wreport_cli,
            ["prepare", "테스트팀",
             "--config", str(project_root / "config.yaml"),
             "--week", "2026-W18",
             "--prev-week", ""],
        )
        assert prep_result.exit_code == 0, f"prepare failed:\n{prep_result.output}\n{prep_result.exception}"
        week_dir = team_data_root / "2026-W18"
        assert (week_dir / "전략기획.xlsx").exists()
        assert (week_dir / "사업개발.xlsx").exists()

        # ----- Step 3: populate 지시사항 -----
        wb = load_workbook(team_data_root / "_지시사항.xlsx")
        ws = wb.active
        # Row 1: 담당파트 column filled → no LLM call expected
        ws.append(["2026-04-27", "신규 분석", "전략기획", "높음", "2026-04-30", ""])
        # Row 2: 담당파트 empty → LLM inference
        ws.append(["2026-04-27", "고객사 미팅", "", "보통", "2026-05-01", ""])
        wb.save(team_data_root / "_지시사항.xlsx")

        # ----- Step 4: wreport assign — LLM returns 사업개발 for the empty row -----
        with patch.object(_llm_mod, "shutil") as fake_shutil, \
             patch.object(_llm_mod, "subprocess") as fake_subproc, \
             patch("smtplib.SMTP") as smtp_cls:
            fake_shutil.which.return_value = "/fake/codex"
            fake_subproc.run.return_value = MagicMock(
                returncode=0,
                stdout=json.dumps(
                    {"추정_담당파트": "사업개발", "확신도": 0.9, "근거": "고객 키워드"},
                    ensure_ascii=False,
                ),
                stderr="",
            )
            smtp_inst = MagicMock()
            smtp_cls.return_value.__enter__.return_value = smtp_inst

            assign_result = runner.invoke(
                wreport_cli,
                ["assign", "테스트팀",
                 "--config", str(project_root / "config.yaml"),
                 "--week", "2026-W18"],
            )
            assert assign_result.exit_code == 0, (
                f"assign failed:\n{assign_result.output}\n{assign_result.exception}"
            )
            assert smtp_inst.send_message.called  # part leads notified

        # Both rows landed in the right workbook with W18-NNN IDs
        wb = load_workbook(week_dir / "전략기획.xlsx")
        rows1 = list(wb["이번주"].iter_rows(min_row=2, values_only=True))
        assert len(rows1) == 1
        assert rows1[0][0].startswith("W18-")        # 업무ID
        assert "신규 분석" in rows1[0][2]            # 업무_지시내용

        wb = load_workbook(week_dir / "사업개발.xlsx")
        rows2 = list(wb["이번주"].iter_rows(min_row=2, values_only=True))
        assert len(rows2) == 1
        assert rows2[0][0].startswith("W18-")
        assert "고객사 미팅" in rows2[0][2]

        # ----- Step 5: part lead fills in 상태 + 처리결과 -----
        for part in ("전략기획", "사업개발"):
            wb = load_workbook(week_dir / f"{part}.xlsx")
            ws = wb["이번주"]
            # Update the only row's 상태 (col D = idx 3) and 처리결과 (col E = idx 4)
            ws.cell(row=2, column=4, value="진행중")
            ws.cell(row=2, column=5, value="작업 진행 중")
            wb.save(week_dir / f"{part}.xlsx")

        # ----- Step 6: wreport compile -----
        fake_compile_result = {
            "팀_종합_요약": "이번 주 사업그룹 두 항목 모두 진행 중.",
            "지난주": {"요약": "지난 주에서 넘어온 항목 없음.", "항목": []},
            "이번주": {
                "요약": "신규 두 건 진행 중.",
                "항목": [
                    {"그룹": "사업그룹", "파트": "전략기획", "업무ID": rows1[0][0],
                     "업무": "신규 분석", "상태": "진행중",
                     "이번주_처리결과": ["작업 진행 중"], "이슈": [],
                     "리스크_지원요청": [], "다음액션": [],
                     "마감": "2026-04-30", "우선순위": "높음"},
                    {"그룹": "사업그룹", "파트": "사업개발", "업무ID": rows2[0][0],
                     "업무": "고객사 미팅", "상태": "진행중",
                     "이번주_처리결과": ["작업 진행 중"], "이슈": [],
                     "리스크_지원요청": [], "다음액션": [],
                     "마감": "2026-05-01", "우선순위": "보통"},
                ],
            },
            "그룹별_요약": {"사업그룹": {"상태": "정상", "요약": "두 항목 정상 진행"}},
            "그룹장_확인필요": [],
            "미작성_파트": [],
            "주의사항": [],
        }
        with patch.object(_llm_mod, "shutil") as fake_shutil, \
             patch.object(_llm_mod, "subprocess") as fake_subproc, \
             patch("smtplib.SMTP") as smtp_cls:
            fake_shutil.which.return_value = "/fake/codex"
            fake_subproc.run.return_value = MagicMock(
                returncode=0,
                stdout=json.dumps(fake_compile_result, ensure_ascii=False),
                stderr="",
            )
            smtp_inst = MagicMock()
            smtp_cls.return_value.__enter__.return_value = smtp_inst

            compile_result = runner.invoke(
                wreport_cli,
                ["compile", "테스트팀",
                 "--config", str(project_root / "config.yaml"),
                 "--week", "2026-W18"],
            )
            assert compile_result.exit_code == 0, (
                f"compile failed:\n{compile_result.output}\n{compile_result.exception}"
            )
            assert smtp_inst.send_message.called

        dashboard_path = week_dir / "_dashboard.html"
        assert dashboard_path.exists(), f"Dashboard not written. Output:\n{compile_result.output}"
        html = dashboard_path.read_text(encoding="utf-8")
        assert "테스트팀" in html
        assert "2026-W18" in html
        assert "이번 주 사업그룹 두 항목 모두 진행 중" in html
        assert "신규 분석" in html
        assert "고객사 미팅" in html
        assert "사업그룹" in html
    finally:
        if str(project_root) in sys.path:
            sys.path.remove(str(project_root))
