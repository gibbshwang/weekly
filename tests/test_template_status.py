"""Tests for status.py — Phase 2 per-part xlsx progress introspection."""
import importlib.util
from pathlib import Path

from scripts.lib.template_render import render_string
from scripts.xlsx_template import generate_part_xlsx
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent


def _load_status(tmp_path: Path):
    tmpl_path = PROJECT_ROOT / "templates" / "src" / "status.py.tmpl"
    rendered = render_string(tmpl_path.read_text(encoding="utf-8"), {})
    out = tmp_path / "status.py"
    out.write_text(rendered, encoding="utf-8")
    spec = importlib.util.spec_from_file_location("status_mod", out)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _make_part_xlsx(week_dir: Path, part: str, *, with_data: bool) -> Path:
    """Generate a part workbook, optionally seeding one row in 이번주."""
    out = week_dir / f"{part}.xlsx"
    generate_part_xlsx(
        out_path=out, team_name="기획팀", group_name="g1", part_name=part, week="2026-W18",
    )
    if with_data:
        wb = load_workbook(out)
        wb["이번주"].append(["W18-001", "지시사항", "task A", "진행중",
                          "결과", "", "", "", "2026-04-30", "보통", ""])
        wb.save(out)
    return out


def test_compute_status_marks_missing(tmp_path: Path):
    s = _load_status(tmp_path)
    week_dir = tmp_path / "week"
    week_dir.mkdir()
    _make_part_xlsx(week_dir, "전략기획", with_data=True)
    statuses = s.compute_status(week_dir, parts=["전략기획", "사업개발"])
    a = next(x for x in statuses if x["part"] == "전략기획")
    b = next(x for x in statuses if x["part"] == "사업개발")
    assert a["status"] == "written"
    assert b["status"] == "missing"
    assert b["mtime"] is None


def test_compute_status_marks_empty_file(tmp_path: Path):
    s = _load_status(tmp_path)
    week_dir = tmp_path / "week"
    week_dir.mkdir()
    _make_part_xlsx(week_dir, "전략기획", with_data=False)
    statuses = s.compute_status(week_dir, parts=["전략기획"])
    assert statuses[0]["status"] == "empty"


def test_compute_status_includes_row_count_and_mtime(tmp_path: Path):
    s = _load_status(tmp_path)
    week_dir = tmp_path / "week"
    week_dir.mkdir()
    out = _make_part_xlsx(week_dir, "전략기획", with_data=True)
    # Add another row
    wb = load_workbook(out)
    wb["이번주"].append(["W18-002", "지시사항", "task B", "미착수",
                      "", "", "", "", "2026-05-01", "낮음", ""])
    wb.save(out)
    statuses = s.compute_status(week_dir, parts=["전략기획"])
    assert statuses[0]["rows"] == 2
    assert statuses[0]["mtime"] is not None


def test_format_status_report_has_all_parts(tmp_path: Path):
    s = _load_status(tmp_path)
    statuses = [
        {"part": "전략기획", "status": "written", "rows": 3, "mtime": "2026-04-25T10:00:00"},
        {"part": "사업개발", "status": "missing", "rows": 0, "mtime": None},
    ]
    text = s.format_status_report(statuses, week="2026-W18", team_name="기획팀")
    assert "기획팀" in text
    assert "2026-W18" in text
    assert "전략기획" in text
    assert "사업개발" in text
    assert "✓" in text
    assert "✗" in text
    assert "미작성" in text


def test_format_status_report_includes_empty_marker(tmp_path: Path):
    s = _load_status(tmp_path)
    statuses = [
        {"part": "전략기획", "status": "empty", "rows": 0, "mtime": "2026-04-25T10:00:00"},
    ]
    text = s.format_status_report(statuses, week="2026-W18", team_name="기획팀")
    assert "○" in text
    assert "비어있음" in text
