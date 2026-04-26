from pathlib import Path
from openpyxl import load_workbook
from scripts.xlsx_template import generate_xlsx


def test_generate_xlsx_has_headers(tmp_path: Path):
    out = tmp_path / "_지시사항.xlsx"
    generate_xlsx(out_path=out, dept_name="기획팀", parts=["전략기획", "사업개발"])
    assert out.exists()
    wb = load_workbook(out)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert headers == ["일자", "지시내용", "담당파트", "우선순위", "마감", "비고"]


def test_generate_xlsx_has_part_dropdown(tmp_path: Path):
    out = tmp_path / "_지시사항.xlsx"
    generate_xlsx(out_path=out, dept_name="기획팀", parts=["전략기획", "사업개발"])
    wb = load_workbook(out)
    ws = wb.active
    # FIX-02: parts now live in the hidden _refs sheet; the DV references that
    # range. The "all parts present" property is now an assertion on _refs cells.
    refs = wb["_refs"]
    cell_values = {refs.cell(row=i, column=1).value for i in range(1, 3)}
    assert {"전략기획", "사업개발"}.issubset(cell_values)
    formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert any("_refs" in (f or "") for f in formulas)


def test_generate_xlsx_has_priority_dropdown(tmp_path: Path):
    out = tmp_path / "_지시사항.xlsx"
    generate_xlsx(out_path=out, dept_name="기획팀", parts=["전략기획"])
    wb = load_workbook(out)
    ws = wb.active
    dvs = ws.data_validations.dataValidation
    formulas = [dv.formula1 for dv in dvs]
    assert any("높음" in f and "보통" in f and "낮음" in f for f in formulas)


def test_generate_xlsx_creates_parent_dir(tmp_path: Path):
    out = tmp_path / "공유폴더" / "기획팀" / "_지시사항.xlsx"
    generate_xlsx(out_path=out, dept_name="기획팀", parts=["전략기획"])
    assert out.exists()


# --- DataValidation hardening (FIX-02) ---


def test_parts_use_hidden_sheet_not_inline_literal(tmp_path: Path):
    """Parts dropdown must reference a hidden sheet, not inline `"a,b,c"`.

    Inline literal sources break on commas/quotes in part names, and
    Excel treats values starting with `=+@-` as formulas if the dropdown
    source is interpreted as a literal list. Sourcing from real cells
    sidesteps both classes of issue.
    """
    out = tmp_path / "_지시사항.xlsx"
    generate_xlsx(out_path=out, dept_name="기획팀", parts=["전략기획", "사업개발"])
    wb = load_workbook(out)
    ws = wb.active
    dv_for_part_col = next(
        dv for dv in ws.data_validations.dataValidation
        if any("C2" in r.coord for r in dv.sqref.ranges)
    )
    formula = dv_for_part_col.formula1 or ""
    assert "_refs" in formula, (
        f"Parts dropdown should reference _refs sheet, got: {formula!r}"
    )
    # Values must NOT be inlined as a literal CSV-in-quotes
    assert "전략기획" not in formula
    assert "사업개발" not in formula


def test_refs_sheet_contains_parts_and_is_hidden(tmp_path: Path):
    """The _refs sheet must hold part names as cell values and be hidden."""
    out = tmp_path / "_지시사항.xlsx"
    generate_xlsx(out_path=out, dept_name="기획팀", parts=["전략기획", "사업개발", "neo-team"])
    wb = load_workbook(out)
    assert "_refs" in wb.sheetnames, f"Expected hidden _refs sheet, got: {wb.sheetnames}"
    refs = wb["_refs"]
    assert refs.sheet_state == "hidden", (
        f"_refs sheet must be hidden, state={refs.sheet_state!r}"
    )
    cell_values = [refs.cell(row=i, column=1).value for i in range(1, 4)]
    assert cell_values == ["전략기획", "사업개발", "neo-team"]


# --- Phase 2 (Wave A.3): generate_directives_xlsx with 담당그룹 column ---

import pytest
from dataclasses import dataclass
from typing import List


@dataclass
class _StubPart:
    name: str


@dataclass
class _StubGroup:
    name: str
    parts: List[_StubPart]


@dataclass
class _StubTeam:
    name: str
    groups: List[_StubGroup]

    def iter_parts(self):
        for g in self.groups:
            for p in g.parts:
                yield g, p


def _stub_team():
    return _StubTeam(
        name="기획팀",
        groups=[
            _StubGroup("사업그룹", [_StubPart("전략기획"), _StubPart("사업개발")]),
            _StubGroup("운영그룹", [_StubPart("운영관리")]),
        ],
    )


def test_directives_xlsx_v2_has_seven_column_header(tmp_path: Path):
    """Phase 2 _지시사항.xlsx adds 담당그룹 column between 지시내용 and 담당파트."""
    from scripts.xlsx_template import generate_directives_xlsx
    out = tmp_path / "_지시사항.xlsx"
    generate_directives_xlsx(out_path=out, team=_stub_team())
    wb = load_workbook(out)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert headers == [
        "일자", "지시내용", "담당그룹", "담당파트", "우선순위", "마감", "비고",
    ]


def test_directives_xlsx_v2_group_dropdown_uses_refs(tmp_path: Path):
    """담당그룹 dropdown sources from _refs hidden sheet (column A holds groups)."""
    from scripts.xlsx_template import generate_directives_xlsx
    out = tmp_path / "_지시사항.xlsx"
    generate_directives_xlsx(out_path=out, team=_stub_team())
    wb = load_workbook(out)
    ws = wb.active
    # 담당그룹 = column C
    dv_group = next(
        dv for dv in ws.data_validations.dataValidation
        if any("C2" in r.coord for r in dv.sqref.ranges)
    )
    assert "_refs" in (dv_group.formula1 or "")
    assert "사업그룹" not in (dv_group.formula1 or "")  # not inlined

    refs = wb["_refs"]
    assert refs.sheet_state == "hidden"
    group_cells = [refs.cell(row=i, column=1).value for i in range(1, 3)]
    assert group_cells == ["사업그룹", "운영그룹"]


def test_directives_xlsx_v2_part_dropdown_lists_all_parts_from_all_groups(tmp_path: Path):
    """담당파트 dropdown holds the flattened part list across all groups."""
    from scripts.xlsx_template import generate_directives_xlsx
    out = tmp_path / "_지시사항.xlsx"
    generate_directives_xlsx(out_path=out, team=_stub_team())
    wb = load_workbook(out)
    ws = wb.active
    # 담당파트 = column D
    dv_part = next(
        dv for dv in ws.data_validations.dataValidation
        if any("D2" in r.coord for r in dv.sqref.ranges)
    )
    assert "_refs" in (dv_part.formula1 or "")

    refs = wb["_refs"]
    part_cells = [refs.cell(row=i, column=2).value for i in range(1, 4)]
    assert part_cells == ["전략기획", "사업개발", "운영관리"]


def test_directives_xlsx_v2_priority_dropdown_unchanged(tmp_path: Path):
    """우선순위 dropdown stays as a controlled inline literal (column E)."""
    from scripts.xlsx_template import generate_directives_xlsx
    out = tmp_path / "_지시사항.xlsx"
    generate_directives_xlsx(out_path=out, team=_stub_team())
    wb = load_workbook(out)
    ws = wb.active
    formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert any("높음" in (f or "") and "낮음" in (f or "") for f in formulas)


# --- Phase 2 (Wave A.3): generate_part_xlsx (per-part workbook) ---


def test_part_xlsx_has_four_sheets(tmp_path: Path):
    """Per-part workbook layout: 지난주 / 이번주 / 지난주_완료 / _refs (hidden)."""
    from scripts.xlsx_template import generate_part_xlsx
    out = tmp_path / "전략기획.xlsx"
    generate_part_xlsx(out_path=out, team_name="기획팀", group_name="사업그룹", part_name="전략기획", week="2026-W18")
    wb = load_workbook(out)
    # All four sheets must exist
    for name in ("지난주", "이번주", "지난주_완료", "_refs"):
        assert name in wb.sheetnames, f"missing sheet {name}; got {wb.sheetnames}"
    assert wb["_refs"].sheet_state == "hidden"


def test_part_xlsx_each_active_sheet_has_eleven_columns(tmp_path: Path):
    """The 지난주, 이번주, 지난주_완료 sheets all share the same 11-column shape."""
    from scripts.xlsx_template import generate_part_xlsx
    out = tmp_path / "전략기획.xlsx"
    generate_part_xlsx(out_path=out, team_name="기획팀", group_name="사업그룹", part_name="전략기획", week="2026-W18")
    wb = load_workbook(out)
    expected_headers = [
        "업무ID", "출처", "업무_지시내용", "상태",
        "이번주_처리결과", "이슈_장애", "리스크_지원요청",
        "다음액션_차주계획", "마감", "우선순위", "비고",
    ]
    for sheet_name in ("지난주", "이번주", "지난주_완료"):
        ws = wb[sheet_name]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert headers == expected_headers, f"{sheet_name} headers mismatch: {headers}"


def test_part_xlsx_has_state_dropdown_on_active_sheets(tmp_path: Path):
    """상태 dropdown values must include the agreed 6-state set on 지난주 + 이번주."""
    from scripts.xlsx_template import generate_part_xlsx
    out = tmp_path / "전략기획.xlsx"
    generate_part_xlsx(out_path=out, team_name="기획팀", group_name="사업그룹", part_name="전략기획", week="2026-W18")
    wb = load_workbook(out)
    expected_states = {"미착수", "진행중", "완료", "지연", "보류", "취소"}
    for sheet_name in ("지난주", "이번주"):
        ws = wb[sheet_name]
        # 상태 = column D (4th)
        dv_state = next(
            dv for dv in ws.data_validations.dataValidation
            if any("D2" in r.coord for r in dv.sqref.ranges)
        )
        formula = dv_state.formula1 or ""
        for state in expected_states:
            assert state in formula, f"{sheet_name} 상태 dropdown missing {state}: {formula}"


def test_part_xlsx_has_source_dropdown(tmp_path: Path):
    """출처 dropdown lists the agreed sources (지시사항 / 파트작성 / 회의 / 고객요청 / 이월 / 기타)."""
    from scripts.xlsx_template import generate_part_xlsx
    out = tmp_path / "전략기획.xlsx"
    generate_part_xlsx(out_path=out, team_name="기획팀", group_name="사업그룹", part_name="전략기획", week="2026-W18")
    wb = load_workbook(out)
    ws = wb["이번주"]
    # 출처 = column B (2nd)
    dv_src = next(
        dv for dv in ws.data_validations.dataValidation
        if any("B2" in r.coord for r in dv.sqref.ranges)
    )
    formula = dv_src.formula1 or ""
    for src in ("지시사항", "파트작성", "회의", "고객요청", "이월", "기타"):
        assert src in formula, f"출처 dropdown missing {src}: {formula}"


def test_part_xlsx_priority_dropdown(tmp_path: Path):
    from scripts.xlsx_template import generate_part_xlsx
    out = tmp_path / "전략기획.xlsx"
    generate_part_xlsx(out_path=out, team_name="기획팀", group_name="사업그룹", part_name="전략기획", week="2026-W18")
    wb = load_workbook(out)
    ws = wb["이번주"]
    formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert any("높음" in (f or "") and "낮음" in (f or "") for f in formulas)


def test_part_xlsx_creates_parent_dir(tmp_path: Path):
    from scripts.xlsx_template import generate_part_xlsx
    out = tmp_path / "공유폴더" / "기획팀" / "2026-W18" / "전략기획.xlsx"
    generate_part_xlsx(out_path=out, team_name="기획팀", group_name="사업그룹", part_name="전략기획", week="2026-W18")
    assert out.exists()


def test_part_xlsx_active_sheet_is_이번주(tmp_path: Path):
    """When the part lead opens the file they should land on '이번주' (this week)."""
    from scripts.xlsx_template import generate_part_xlsx
    out = tmp_path / "전략기획.xlsx"
    generate_part_xlsx(out_path=out, team_name="기획팀", group_name="사업그룹", part_name="전략기획", week="2026-W18")
    wb = load_workbook(out)
    assert wb.active.title == "이번주"
