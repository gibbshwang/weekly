"""Tests for the Phase 2 assigner.py.tmpl — xlsx-based assign with 업무ID."""
import importlib.util
import json
from dataclasses import dataclass
from pathlib import Path
from typing import List
from unittest.mock import MagicMock

from openpyxl import load_workbook

from scripts.lib.template_render import render_string
from scripts.xlsx_template import generate_part_xlsx, PART_SHEET_HEADERS


PROJECT_ROOT = Path(__file__).resolve().parent.parent
PROMPT_PATH = PROJECT_ROOT / "templates/src/prompts/assign_system.txt.tmpl"


def _load_assigner(tmp_path: Path):
    text = render_string(
        (PROJECT_ROOT / "templates/src/assigner.py.tmpl").read_text(encoding="utf-8"), {}
    )
    out = tmp_path / "assigner.py"
    out.write_text(text, encoding="utf-8")
    spec = importlib.util.spec_from_file_location("assigner_mod", out)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _state_path(tmp_path: Path) -> Path:
    return tmp_path / "_assignments_state.json"


@dataclass
class _StubPart:
    name: str
    role: "str | None" = None


@dataclass
class _StubGroup:
    name: str
    parts: List[_StubPart]


@dataclass
class _StubTeam:
    name: str
    groups: List[_StubGroup]

    def iter_parts(self):
        for g in self.groups:
            for p in g.parts:
                yield g, p


def _team():
    return _StubTeam(
        name="기획팀",
        groups=[
            _StubGroup("사업그룹", [_StubPart("전략기획"), _StubPart("사업개발")]),
            _StubGroup("운영그룹", [_StubPart("운영관리")]),
        ],
    )


def _create_week_workbooks(week_dir: Path, team) -> None:
    week_dir.mkdir(parents=True, exist_ok=True)
    for g, p in team.iter_parts():
        generate_part_xlsx(
            out_path=week_dir / f"{p.name}.xlsx",
            team_name=team.name, group_name=g.name,
            part_name=p.name, week="2026-W18",
        )


# --- instruction_key + diff_rows + state plumbing (Todo 2: content fingerprint) ---


def test_instruction_key_is_position_independent(tmp_path: Path):
    """Same content at different row positions yields the same key — that's
    the whole point of the fingerprint vs row_index migration."""
    a = _load_assigner(tmp_path)
    row = {"일자": "2026-04-22", "지시내용": "x", "담당파트": "전략기획",
           "우선순위": "높음", "마감": "2026-04-30", "비고": "anything"}
    assert a._instruction_key(row) == a._instruction_key(dict(row))


def test_instruction_key_changes_when_지시내용_changes(tmp_path: Path):
    a = _load_assigner(tmp_path)
    base = {"일자": "2026-04-22", "지시내용": "원본", "담당파트": "전략기획",
            "우선순위": "높음", "마감": "2026-04-30"}
    edited = {**base, "지시내용": "수정됨"}
    assert a._instruction_key(base) != a._instruction_key(edited)


def test_instruction_key_ignores_비고(tmp_path: Path):
    """비고 is operator commentary, not part of the instruction identity.
    Editing 비고 must NOT trigger re-assignment."""
    a = _load_assigner(tmp_path)
    base = {"일자": "2026-04-22", "지시내용": "x", "담당파트": "전략기획",
            "우선순위": "높음", "마감": "2026-04-30", "비고": ""}
    note_added = {**base, "비고": "임원 보고용"}
    assert a._instruction_key(base) == a._instruction_key(note_added)


def test_instruction_key_normalizes_whitespace(tmp_path: Path):
    """Trailing/leading ws on operator-edited cells should not flip identity."""
    a = _load_assigner(tmp_path)
    a_row = {"일자": "2026-04-22", "지시내용": "x", "담당파트": "전략기획",
             "우선순위": "높음", "마감": "2026-04-30"}
    b_row = {"일자": "2026-04-22 ", "지시내용": " x ", "담당파트": "전략기획",
             "우선순위": "높음", "마감": "  2026-04-30"}
    assert a._instruction_key(a_row) == a._instruction_key(b_row)


def test_diff_rows_returns_indices_of_unseen_keys(tmp_path: Path):
    """Empty state → all rows are new."""
    a = _load_assigner(tmp_path)
    rows = [
        {"일자": "2026-04-22", "지시내용": "x", "담당파트": "전략기획",
         "우선순위": "높음", "마감": "2026-04-30", "비고": ""},
    ]
    assert a.diff_rows(rows, _state_path(tmp_path)) == [0]


def test_diff_rows_skips_known_keys(tmp_path: Path):
    a = _load_assigner(tmp_path)
    sp = _state_path(tmp_path)
    rows = [
        {"일자": "2026-04-22", "지시내용": "x", "담당파트": "전략기획",
         "우선순위": "높음", "마감": "2026-04-30", "비고": ""},
    ]
    a.update_state(sp, rows, [(0, "전략기획", "column")])
    assert a.diff_rows(rows, sp) == []


def test_diff_rows_resilient_to_row_reorder(tmp_path: Path):
    """Re-sorting the directives xlsx must NOT cause re-assignment.
    This is the core bug Todo 2 is fixing."""
    a = _load_assigner(tmp_path)
    sp = _state_path(tmp_path)
    rows_v1 = [
        {"일자": "2026-04-22", "지시내용": "first", "담당파트": "전략기획",
         "우선순위": "높음", "마감": "2026-04-30", "비고": ""},
        {"일자": "2026-04-23", "지시내용": "second", "담당파트": "사업개발",
         "우선순위": "보통", "마감": "2026-05-02", "비고": ""},
    ]
    a.update_state(sp, rows_v1, [(0, "전략기획", "column"), (1, "사업개발", "column")])
    # Sorted/swapped — same content, different positions
    rows_v2 = [rows_v1[1], rows_v1[0]]
    assert a.diff_rows(rows_v2, sp) == []


def test_diff_rows_resilient_to_row_insertion(tmp_path: Path):
    """Inserting a new row at the top must NOT re-classify rows below it."""
    a = _load_assigner(tmp_path)
    sp = _state_path(tmp_path)
    rows_v1 = [
        {"일자": "2026-04-22", "지시내용": "old1", "담당파트": "전략기획",
         "우선순위": "높음", "마감": "2026-04-30", "비고": ""},
    ]
    a.update_state(sp, rows_v1, [(0, "전략기획", "column")])
    rows_v2 = [
        {"일자": "2026-04-21", "지시내용": "NEW", "담당파트": "전략기획",
         "우선순위": "보통", "마감": "2026-04-29", "비고": ""},
        rows_v1[0],   # original — now at index 1
    ]
    new = a.diff_rows(rows_v2, sp)
    assert new == [0]   # only the new top row, NOT the shifted original


def test_diff_rows_treats_content_edit_as_new(tmp_path: Path):
    """If the operator edits 지시내용 in place, the new content has a new
    instruction_key, so it's treated as a new instruction (and gets a new
    task ID). The old key remains in state — out of scope for this PR
    (operator workflow expects edits to be rare)."""
    a = _load_assigner(tmp_path)
    sp = _state_path(tmp_path)
    rows_v1 = [
        {"일자": "2026-04-22", "지시내용": "원본", "담당파트": "전략기획",
         "우선순위": "높음", "마감": "2026-04-30", "비고": ""},
    ]
    a.update_state(sp, rows_v1, [(0, "전략기획", "column")])
    rows_v2 = [{**rows_v1[0], "지시내용": "수정됨"}]
    assert a.diff_rows(rows_v2, sp) == [0]


def test_state_uses_instructions_schema(tmp_path: Path):
    """State on disk must have an `instructions` array (new schema), not
    the legacy `rows` array."""
    a = _load_assigner(tmp_path)
    sp = _state_path(tmp_path)
    rows = [
        {"일자": "2026-04-22", "지시내용": "x", "담당파트": "전략기획",
         "우선순위": "높음", "마감": "2026-04-30", "비고": ""},
    ]
    a.update_state(sp, rows, [(0, "전략기획", "column")])
    payload = json.loads(sp.read_text(encoding="utf-8"))
    assert "instructions" in payload
    assert "rows" not in payload   # legacy schema fully removed
    e0 = payload["instructions"][0]
    assert "instruction_key" in e0
    assert e0["assigned_to"] == "전략기획"
    assert e0["method"] == "column"


# --- generate_task_id + next_seq_for_week ---


def test_generate_task_id_w_format(tmp_path: Path):
    a = _load_assigner(tmp_path)
    assert a.generate_task_id("2026-W18", 3) == "W18-003"
    assert a.generate_task_id("2026-W01", 1) == "W01-001"
    assert a.generate_task_id("2026-W52", 999) == "W52-999"


def test_next_seq_for_week_starts_at_one_when_empty(tmp_path: Path):
    a = _load_assigner(tmp_path)
    week_dir = tmp_path / "2026-W18"
    _create_week_workbooks(week_dir, _team())
    assert a.next_seq_for_week(week_dir, _team(), week="2026-W18") == 1


def test_next_seq_for_week_continues_from_max(tmp_path: Path):
    a = _load_assigner(tmp_path)
    team = _team()
    week_dir = tmp_path / "2026-W18"
    _create_week_workbooks(week_dir, team)
    wb1 = load_workbook(week_dir / "전략기획.xlsx")
    wb1["이번주"].append(["W18-005"] + ["" for _ in PART_SHEET_HEADERS[1:]])
    wb1.save(week_dir / "전략기획.xlsx")
    wb2 = load_workbook(week_dir / "운영관리.xlsx")
    wb2["이번주"].append(["W18-007"] + ["" for _ in PART_SHEET_HEADERS[1:]])
    wb2.save(week_dir / "운영관리.xlsx")
    assert a.next_seq_for_week(week_dir, team, week="2026-W18") == 8


def test_next_seq_for_week_ignores_other_week_ids(tmp_path: Path):
    a = _load_assigner(tmp_path)
    team = _team()
    week_dir = tmp_path / "2026-W18"
    _create_week_workbooks(week_dir, team)
    wb = load_workbook(week_dir / "전략기획.xlsx")
    wb["지난주"].append(["W17-099"] + ["" for _ in PART_SHEET_HEADERS[1:]])
    wb.save(week_dir / "전략기획.xlsx")
    assert a.next_seq_for_week(week_dir, team, week="2026-W18") == 1


# --- resolve_target_part (group-aware) ---


def test_resolve_target_part_uses_column_when_filled(tmp_path: Path):
    a = _load_assigner(tmp_path)
    fake_llm = MagicMock()
    team = _team()
    row = {"일자": "x", "지시내용": "x",
           "담당파트": "전략기획", "우선순위": "보통", "마감": "x", "비고": ""}
    group, part, method = a.resolve_target_part(
        row, team=team, llm=fake_llm,
        prompt_template_path=PROMPT_PATH,
    )
    assert part.name == "전략기획"
    assert group.name == "사업그룹"
    assert method == "column"
    fake_llm.call.assert_not_called()


def test_resolve_target_part_calls_llm_when_part_empty(tmp_path: Path):
    a = _load_assigner(tmp_path)
    fake_llm = MagicMock()
    fake_llm.call.return_value = json.dumps({
        "추정_담당파트": "사업개발", "확신도": 0.8, "근거": "고객사 미팅"
    }, ensure_ascii=False)
    team = _team()
    row = {"일자": "x", "지시내용": "고객사 미팅", "담당파트": "", "우선순위": "x", "마감": "x", "비고": ""}
    group, part, method = a.resolve_target_part(
        row, team=team, llm=fake_llm,
        prompt_template_path=PROMPT_PATH,
    )
    assert part.name == "사업개발"
    assert group.name == "사업그룹"
    assert method == "ai_inferred"


def test_resolve_target_part_falls_back_when_llm_returns_invalid(tmp_path: Path):
    a = _load_assigner(tmp_path)
    fake_llm = MagicMock()
    fake_llm.call.return_value = json.dumps({
        "추정_담당파트": "무효한파트", "확신도": 0.3, "근거": "x"
    }, ensure_ascii=False)
    team = _team()
    row = {"일자": "x", "지시내용": "x", "담당파트": "", "우선순위": "x", "마감": "x", "비고": ""}
    group, part, method = a.resolve_target_part(
        row, team=team, llm=fake_llm, prompt_template_path=PROMPT_PATH,
    )
    # Falls back to first (group, part) of team
    assert part.name == "전략기획"
    assert group.name == "사업그룹"
    assert method == "ai_inferred"


def test_resolve_target_part_handles_codex_code_fence(tmp_path: Path):
    a = _load_assigner(tmp_path)
    fake_llm = MagicMock()
    fake_llm.call.return_value = '```json\n{"추정_담당파트":"전략기획","확신도":0.7,"근거":"전략 키워드"}\n```'
    row = {"일자": "x", "지시내용": "전략 수립", "담당파트": "", "우선순위": "x", "마감": "x", "비고": ""}
    group, part, method = a.resolve_target_part(
        row, team=_team(), llm=fake_llm, prompt_template_path=PROMPT_PATH,
    )
    assert part.name == "전략기획"


# --- Part.role propagation into LLM prompt (Todo 1: role 마스터) ---


def _team_with_roles():
    return _StubTeam(
        name="기획팀",
        groups=[
            _StubGroup("사업그룹", [
                _StubPart("전략기획", role="중장기 전략, 시장 분석, 경쟁사 조사"),
                _StubPart("사업개발", role="고객사 미팅, 신규 사업 발굴"),
            ]),
            _StubGroup("운영그룹", [
                _StubPart("운영관리"),  # role intentionally omitted
            ]),
        ],
    )


def test_resolve_target_part_includes_part_role_in_prompt(tmp_path: Path):
    """When parts have `role` set, the LLM system prompt must surface them so
    the model can reason about which part owns the instruction."""
    a = _load_assigner(tmp_path)
    fake_llm = MagicMock()
    fake_llm.call.return_value = json.dumps({
        "추정_담당파트": "사업개발", "확신도": 0.9, "근거": "고객사 키워드"
    }, ensure_ascii=False)
    team = _team_with_roles()
    row = {"일자": "x", "지시내용": "고객사 미팅", "담당파트": "", "우선순위": "x", "마감": "x", "비고": ""}
    a.resolve_target_part(
        row, team=team, llm=fake_llm, prompt_template_path=PROMPT_PATH,
    )
    sys_prompt = fake_llm.call.call_args[0][0]
    assert "전략기획" in sys_prompt
    assert "중장기 전략, 시장 분석, 경쟁사 조사" in sys_prompt
    assert "사업개발" in sys_prompt
    assert "고객사 미팅, 신규 사업 발굴" in sys_prompt


def test_resolve_target_part_handles_parts_without_role(tmp_path: Path):
    """A part with role=None must still appear in the prompt, and each part
    should be on its own line so a role-less part is unambiguously identifiable
    (no role descriptor appended)."""
    a = _load_assigner(tmp_path)
    fake_llm = MagicMock()
    fake_llm.call.return_value = json.dumps({
        "추정_담당파트": "운영관리", "확신도": 0.6, "근거": "운영 키워드"
    }, ensure_ascii=False)
    team = _team_with_roles()
    row = {"일자": "x", "지시내용": "운영 점검", "담당파트": "", "우선순위": "x", "마감": "x", "비고": ""}
    a.resolve_target_part(
        row, team=team, llm=fake_llm, prompt_template_path=PROMPT_PATH,
    )
    sys_prompt = fake_llm.call.call_args[0][0]
    # 운영관리 must be on its own line (proves per-part rendering, not the
    # legacy one-line "파트 목록: a, b, c" format).
    lines = sys_prompt.splitlines()
    운영_only_lines = [l for l in lines if "운영관리" in l and "전략기획" not in l and "사업개발" not in l]
    assert len(운영_only_lines) == 1, (
        f"운영관리 should appear on its own per-part line, found: {운영_only_lines!r}"
    )
    # And that line must NOT carry a role descriptor (role=None for this part).
    assert "—" not in 운영_only_lines[0], (
        "role-less part must not get a role descriptor"
    )


# --- append_to_part_xlsx ---


def test_append_to_part_xlsx_neutralizes_formula_in_free_text(tmp_path: Path):
    """Regression: assigner copies free-text fields (지시내용, 비고) from
    _지시사항.xlsx straight into the part workbook. Formula-leading values
    like `=HYPERLINK(..)` would auto-evaluate when the part lead opens
    their xlsx. safe_excel_text must prefix `'` to neutralize."""
    a = _load_assigner(tmp_path)
    team = _team()
    week_dir = tmp_path / "2026-W18"
    _create_week_workbooks(week_dir, team)
    a.append_to_part_xlsx(
        part_xlsx_path=week_dir / "전략기획.xlsx",
        task_id="W18-001",
        row={
            "일자": "2026-04-27",
            "지시내용": "=HYPERLINK(\"https://evil/\",\"클릭\")",
            "우선순위": "높음",
            "마감": "2026-04-30",
            "비고": "+cmd|' /C calc'!A0",
        },
        source="지시사항",
    )
    wb = load_workbook(week_dir / "전략기획.xlsx")
    rows = list(wb["이번주"].iter_rows(min_row=2, values_only=True))
    by_header = dict(zip(PART_SHEET_HEADERS, rows[0]))
    # Both free-text cells start with `'` (Excel renders as plain text)
    assert by_header["업무_지시내용"].startswith("'=HYPERLINK"), by_header["업무_지시내용"]
    assert by_header["비고"].startswith("'+cmd|"), by_header["비고"]
    # System-generated cols (업무ID, 출처) untouched
    assert by_header["업무ID"] == "W18-001"
    assert by_header["출처"] == "지시사항"


def test_append_to_part_xlsx_writes_to_이번주_with_source_지시사항(tmp_path: Path):
    a = _load_assigner(tmp_path)
    team = _team()
    week_dir = tmp_path / "2026-W18"
    _create_week_workbooks(week_dir, team)
    a.append_to_part_xlsx(
        part_xlsx_path=week_dir / "전략기획.xlsx",
        task_id="W18-001",
        row={"일자": "2026-04-27", "지시내용": "신규 분석",
             "우선순위": "높음", "마감": "2026-04-30", "비고": "임원 보고용"},
        source="지시사항",
    )
    wb = load_workbook(week_dir / "전략기획.xlsx")
    rows = list(wb["이번주"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    by_header = dict(zip(PART_SHEET_HEADERS, rows[0]))
    assert by_header["업무ID"] == "W18-001"
    assert by_header["출처"] == "지시사항"
    assert by_header["업무_지시내용"] == "신규 분석"
    assert by_header["마감"] == "2026-04-30"
    assert by_header["우선순위"] == "높음"
    assert by_header["비고"] == "임원 보고용"


# --- run_assign end-to-end ---


def test_run_assign_appends_to_correct_parts_with_task_ids(tmp_path: Path):
    a = _load_assigner(tmp_path)
    fake_llm = MagicMock()
    team = _team()
    week_dir = tmp_path / "2026-W18"
    _create_week_workbooks(week_dir, team)

    rows = [
        {"일자": "2026-04-27", "지시내용": "전략 task",
         "담당파트": "전략기획", "우선순위": "높음", "마감": "2026-04-30", "비고": ""},
        {"일자": "2026-04-27", "지시내용": "운영 task",
         "담당파트": "운영관리", "우선순위": "보통", "마감": "2026-05-01", "비고": ""},
    ]
    affected = a.run_assign(
        rows=rows, team=team, week="2026-W18",
        state_path=tmp_path / "_state.json",
        part_xlsx_dir=week_dir,
        llm=fake_llm,
        prompt_template_path=PROMPT_PATH,
    )
    assert sorted(affected.keys()) == ["운영관리", "전략기획"]
    all_ids = affected["전략기획"] + affected["운영관리"]
    assert "W18-001" in all_ids
    assert "W18-002" in all_ids

    rows1 = list(load_workbook(week_dir / "전략기획.xlsx")["이번주"].iter_rows(min_row=2, values_only=True))
    assert len(rows1) == 1
    assert "전략 task" in rows1[0]
    rows2 = list(load_workbook(week_dir / "운영관리.xlsx")["이번주"].iter_rows(min_row=2, values_only=True))
    assert len(rows2) == 1
    assert "운영 task" in rows2[0]


def test_run_assign_raises_clear_error_when_part_workbooks_missing(tmp_path: Path):
    """Regression: if `wreport prepare <팀>` was skipped (or week_dir wasn't
    created), assign should fail with a message that points the operator at
    the missing step, not a confusing low-level openpyxl FileNotFoundError."""
    a = _load_assigner(tmp_path)
    fake_llm = MagicMock()
    team = _team()
    week_dir = tmp_path / "2026-W18"  # NOT created — prepare was skipped
    state_path = tmp_path / "_state.json"

    rows = [
        {"일자": "x", "지시내용": "task A",
         "담당파트": "전략기획", "우선순위": "높음", "마감": "x", "비고": ""},
    ]
    import pytest
    with pytest.raises(FileNotFoundError) as excinfo:
        a.run_assign(rows=rows, team=team, week="2026-W18",
                     state_path=state_path, part_xlsx_dir=week_dir,
                     llm=fake_llm, prompt_template_path=PROMPT_PATH)
    msg = str(excinfo.value)
    assert "prepare" in msg, f"error must mention 'prepare' to guide the operator: {msg}"
    assert "2026-W18" in msg, f"error must include the week: {msg}"


def test_run_assign_state_in_team_root_prevents_cross_week_reassign(tmp_path: Path):
    """Regression for the CRITICAL multi-week duplicate bug:

    _지시사항.xlsx is at team root and is long-lived (operators don't clear it
    between weeks). assign state is keyed by instruction_key, but if state
    is stored *inside* the week_dir, every new week starts with empty state
    and re-assigns every row that's still in 지시사항.xlsx.

    Combined with prepare's carry-forward (prior week's incomplete rows go
    into 지난주 sheet), this produces:
      - 지난주 sheet: 이월 row from prepare
      - 이번주 sheet: SAME instruction re-assigned by assign

    Putting state at team root (shared across weeks) prevents this.
    """
    a = _load_assigner(tmp_path)
    fake_llm = MagicMock()
    team = _team()
    team_root = tmp_path / "기획팀"
    team_root.mkdir(parents=True)

    # Single team-root state file (the fix)
    state_path = team_root / "_assignments_state.json"

    # Week 17 setup + assign
    week17_dir = team_root / "2026-W17"
    _create_week_workbooks(week17_dir, team)
    rows = [
        {"일자": "2026-04-22", "지시내용": "long-running task",
         "담당파트": "전략기획", "우선순위": "높음", "마감": "2026-04-30", "비고": ""},
    ]
    affected_w17 = a.run_assign(
        rows=rows, team=team, week="2026-W17",
        state_path=state_path, part_xlsx_dir=week17_dir,
        llm=fake_llm, prompt_template_path=PROMPT_PATH,
    )
    assert "전략기획" in affected_w17

    # Now week 18 — fresh week_dir, but the SAME state_path at team root.
    # Operator hasn't removed the row from 지시사항.xlsx yet.
    week18_dir = team_root / "2026-W18"
    _create_week_workbooks(week18_dir, team)
    affected_w18 = a.run_assign(
        rows=rows, team=team, week="2026-W18",
        state_path=state_path, part_xlsx_dir=week18_dir,
        llm=fake_llm, prompt_template_path=PROMPT_PATH,
    )
    # The fingerprint is in shared state → no re-assignment
    assert affected_w18 == {}, (
        f"team-root state_path must prevent cross-week re-assign; got {affected_w18!r}"
    )
    # Week 18 workbook has NO duplicate row in 이번주
    rows_w18 = list(load_workbook(week18_dir / "전략기획.xlsx")["이번주"]
                    .iter_rows(min_row=2, values_only=True))
    assert rows_w18 == [], f"이번주 should be empty for w18, got {rows_w18}"


def test_run_assign_skips_when_lockfile_held_by_concurrent_run(tmp_path: Path):
    """Regression: hourly assign cron + LLM fallback / SMTP delay can let
    two ticks overlap. Without a lock both processes append to the same
    part workbook → duplicate rows in 이번주.

    run_assign acquires a sidecar lockfile beside state_path. If another
    process holds it (and the lock isn't stale), this tick must skip
    silently — return {} and not modify any workbook."""
    a = _load_assigner(tmp_path)
    fake_llm = MagicMock()
    team = _team()
    team_root = tmp_path / "기획팀"
    team_root.mkdir()
    state_path = team_root / "_assignments_state.json"
    week_dir = team_root / "2026-W18"
    _create_week_workbooks(week_dir, team)

    # Pre-create a fresh lockfile (simulating an in-progress concurrent assign)
    lock_path = Path(str(state_path) + ".lock")
    lock_path.write_text(str(99999), encoding="utf-8")  # bogus pid
    import os, time
    # mtime is now → not stale
    os.utime(lock_path, (time.time(), time.time()))

    rows = [
        {"일자": "2026-04-22", "지시내용": "task A",
         "담당파트": "전략기획", "우선순위": "높음", "마감": "2026-04-30", "비고": ""},
    ]
    affected = a.run_assign(
        rows=rows, team=team, week="2026-W18",
        state_path=state_path, part_xlsx_dir=week_dir,
        llm=fake_llm, prompt_template_path=PROMPT_PATH,
    )
    # Locked tick → no work performed
    assert affected == {}
    # Lockfile preserved for the other process
    assert lock_path.exists()
    # Part workbook untouched
    rows_w = list(load_workbook(week_dir / "전략기획.xlsx")["이번주"]
                  .iter_rows(min_row=2, values_only=True))
    assert rows_w == []


def test_run_assign_recovers_stale_lockfile(tmp_path: Path):
    """If a previous assign crashed mid-run, its lockfile may linger.
    Don't deadlock subsequent ticks — recover after the stale threshold."""
    a = _load_assigner(tmp_path)
    fake_llm = MagicMock()
    team = _team()
    team_root = tmp_path / "기획팀"
    team_root.mkdir()
    state_path = team_root / "_assignments_state.json"
    week_dir = team_root / "2026-W18"
    _create_week_workbooks(week_dir, team)

    lock_path = Path(str(state_path) + ".lock")
    lock_path.write_text(str(99999), encoding="utf-8")
    # Make the lock look 1 day old → stale
    import os, time
    one_day_ago = time.time() - 86400
    os.utime(lock_path, (one_day_ago, one_day_ago))

    rows = [
        {"일자": "2026-04-22", "지시내용": "task A",
         "담당파트": "전략기획", "우선순위": "높음", "마감": "2026-04-30", "비고": ""},
    ]
    affected = a.run_assign(
        rows=rows, team=team, week="2026-W18",
        state_path=state_path, part_xlsx_dir=week_dir,
        llm=fake_llm, prompt_template_path=PROMPT_PATH,
    )
    # Stale lock recovered; assign proceeds normally
    assert "전략기획" in affected
    # Lockfile released after run
    assert not lock_path.exists()


def test_run_assign_idempotent_on_second_call(tmp_path: Path):
    a = _load_assigner(tmp_path)
    fake_llm = MagicMock()
    team = _team()
    week_dir = tmp_path / "2026-W18"
    _create_week_workbooks(week_dir, team)
    state_path = tmp_path / "_state.json"

    rows = [
        {"일자": "x", "지시내용": "task A",
         "담당파트": "전략기획", "우선순위": "높음", "마감": "x", "비고": ""},
    ]
    a.run_assign(rows=rows, team=team, week="2026-W18",
                 state_path=state_path, part_xlsx_dir=week_dir,
                 llm=fake_llm, prompt_template_path=PROMPT_PATH)
    affected2 = a.run_assign(rows=rows, team=team, week="2026-W18",
                             state_path=state_path, part_xlsx_dir=week_dir,
                             llm=fake_llm, prompt_template_path=PROMPT_PATH)
    assert affected2 == {}
    rows_sheet = list(load_workbook(week_dir / "전략기획.xlsx")["이번주"].iter_rows(min_row=2, values_only=True))
    assert len(rows_sheet) == 1
