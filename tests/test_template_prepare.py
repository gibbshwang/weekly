"""Tests for templates/src/prepare.py.tmpl — Phase 2 weekly initialization.

prepare is invoked at the start of each weekly cycle (typically by
`wreport prepare <팀명>` on Monday morning, via cron). It:

  1. Creates `<storage_root>/<team>/<week>/` if missing
  2. Generates a per-part xlsx for every (group, part) pair
  3. If a previous week is supplied, triages each prior part workbook:
     - Rows whose 상태 ∈ {완료, 취소}  → moved to that week's 지난주_완료 sheet
     - Rows whose 상태 ∉ {완료, 취소}  → carried to the new week's 지난주 sheet
                                       with 출처 rewritten to "이월"
  4. The original 업무ID is preserved across carry-forward.
"""
import importlib.util
from dataclasses import dataclass
from pathlib import Path
from typing import List

from openpyxl import load_workbook
from scripts.lib.template_render import render_string
from scripts.xlsx_template import generate_part_xlsx, PART_SHEET_HEADERS


PROJECT_ROOT = Path(__file__).resolve().parent.parent


def _load_prepare(tmp_path: Path):
    text = render_string(
        (PROJECT_ROOT / "templates/src/prepare.py.tmpl").read_text(encoding="utf-8"), {}
    )
    out = tmp_path / "prepare.py"
    out.write_text(text, encoding="utf-8")
    spec = importlib.util.spec_from_file_location("prepare", out)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# --- Stubs that match the Pydantic Team / Group / Part shape ---


@dataclass
class _StubPart:
    name: str


@dataclass
class _StubGroup:
    name: str
    parts: List[_StubPart]


@dataclass
class _StubTeam:
    name: str
    groups: List[_StubGroup]

    def iter_parts(self):
        for g in self.groups:
            for p in g.parts:
                yield g, p


def _team():
    return _StubTeam(
        name="기획팀",
        groups=[
            _StubGroup("사업그룹", [_StubPart("전략기획"), _StubPart("사업개발")]),
            _StubGroup("운영그룹", [_StubPart("운영관리")]),
        ],
    )


def _seed_prev_week_part(
    storage_root: Path, team_name: str, prev_week: str,
    group_name: str, part_name: str,
    rows_이번주: List[dict], rows_지난주: List[dict] | None = None,
) -> Path:
    """Build a previous-week part xlsx with caller-provided rows."""
    out = storage_root / team_name / prev_week / f"{part_name}.xlsx"
    generate_part_xlsx(
        out_path=out, team_name=team_name, group_name=group_name,
        part_name=part_name, week=prev_week,
    )
    wb = load_workbook(out)
    for sheet, rows in (("이번주", rows_이번주), ("지난주", rows_지난주 or [])):
        ws = wb[sheet]
        for row in rows:
            ws.append([row.get(h, "") for h in PART_SHEET_HEADERS])
    wb.save(out)
    return out


# --- Tests ---


def test_prepare_creates_week_folder_and_one_xlsx_per_part(tmp_path: Path):
    prep = _load_prepare(tmp_path)
    team = _team()
    storage_root = tmp_path / "store"
    summary = prep.prepare_new_week(
        storage_root=storage_root, team=team, current_week="2026-W18",
    )
    week_dir = storage_root / team.name / "2026-W18"
    assert week_dir.is_dir()
    for _, part in team.iter_parts():
        assert (week_dir / f"{part.name}.xlsx").exists(), (
            f"missing xlsx for part {part.name}"
        )
    # Summary should report which parts were touched
    assert summary["created"] == ["전략기획", "사업개발", "운영관리"]


def test_prepare_first_week_carry_forward_is_empty(tmp_path: Path):
    """New team's first week: prev_week=None → carry-forward sheets stay empty."""
    prep = _load_prepare(tmp_path)
    team = _team()
    storage_root = tmp_path / "store"
    prep.prepare_new_week(
        storage_root=storage_root, team=team,
        current_week="2026-W18", prev_week=None,
    )
    week_dir = storage_root / team.name / "2026-W18"
    for _, part in team.iter_parts():
        wb = load_workbook(week_dir / f"{part.name}.xlsx")
        last = wb["지난주"]
        # Header row only — no data rows
        assert last.max_row == 1


def test_prepare_carries_forward_incomplete_rows(tmp_path: Path):
    """Rows whose 상태 ∉ {완료, 취소} land in the new week's 지난주 sheet,
    with 출처 rewritten to "이월" and 업무ID preserved."""
    prep = _load_prepare(tmp_path)
    team = _team()
    storage_root = tmp_path / "store"

    _seed_prev_week_part(
        storage_root, team.name, "2026-W17",
        group_name="사업그룹", part_name="전략기획",
        rows_이번주=[
            {"업무ID": "W17-001", "출처": "지시사항", "업무_지시내용": "Task A",
             "상태": "진행중", "마감": "2026-04-30", "우선순위": "높음"},
            {"업무ID": "W17-002", "출처": "파트작성", "업무_지시내용": "Task B",
             "상태": "지연", "마감": "2026-05-03", "우선순위": "보통"},
            {"업무ID": "W17-003", "출처": "지시사항", "업무_지시내용": "Task C (done)",
             "상태": "완료", "마감": "2026-04-25", "우선순위": "보통"},
        ],
    )

    prep.prepare_new_week(
        storage_root=storage_root, team=team,
        current_week="2026-W18", prev_week="2026-W17",
    )

    new_xlsx = storage_root / team.name / "2026-W18" / "전략기획.xlsx"
    wb = load_workbook(new_xlsx)
    last = wb["지난주"]
    rows = list(last.iter_rows(min_row=2, values_only=True))
    # Two incomplete rows carried forward; the 완료 one is NOT here
    assert len(rows) == 2
    # Each row preserves 업무ID and rewrites 출처 to "이월"
    by_id = {r[0]: r for r in rows}
    assert "W17-001" in by_id and "W17-002" in by_id
    assert by_id["W17-001"][1] == "이월"
    assert by_id["W17-002"][1] == "이월"
    # Original 업무_지시내용 preserved
    assert "Task A" in by_id["W17-001"][2]


def test_prepare_archives_completed_rows_to_prev_week(tmp_path: Path):
    """Completed/취소 rows from prev week are MOVED to that week's 지난주_완료
    sheet (and removed from 이번주 + 지난주). They are NOT carried forward."""
    prep = _load_prepare(tmp_path)
    team = _team()
    storage_root = tmp_path / "store"

    prev_xlsx = _seed_prev_week_part(
        storage_root, team.name, "2026-W17",
        group_name="사업그룹", part_name="전략기획",
        rows_이번주=[
            {"업무ID": "W17-005", "출처": "지시사항", "업무_지시내용": "Done task",
             "상태": "완료", "마감": "2026-04-25", "우선순위": "보통"},
            {"업무ID": "W17-006", "출처": "파트작성", "업무_지시내용": "Cancelled",
             "상태": "취소", "마감": "2026-04-25", "우선순위": "낮음"},
            {"업무ID": "W17-007", "출처": "지시사항", "업무_지시내용": "Still going",
             "상태": "진행중", "마감": "2026-05-01", "우선순위": "높음"},
        ],
    )

    prep.prepare_new_week(
        storage_root=storage_root, team=team,
        current_week="2026-W18", prev_week="2026-W17",
    )

    # Prev week file: archive populated, 이번주 trimmed
    wb_prev = load_workbook(prev_xlsx)
    archive = wb_prev["지난주_완료"]
    archive_ids = [r[0] for r in archive.iter_rows(min_row=2, values_only=True)]
    assert sorted(archive_ids) == ["W17-005", "W17-006"]

    this_week_prev = wb_prev["이번주"]
    remaining_ids = [r[0] for r in this_week_prev.iter_rows(min_row=2, values_only=True) if r[0]]
    assert remaining_ids == ["W17-007"], (
        f"이번주 should have only the still-incomplete row left, got {remaining_ids}"
    )

    # New week: 지난주 holds the carried-forward incomplete row
    wb_new = load_workbook(storage_root / team.name / "2026-W18" / "전략기획.xlsx")
    new_last = wb_new["지난주"]
    new_ids = [r[0] for r in new_last.iter_rows(min_row=2, values_only=True) if r[0]]
    assert new_ids == ["W17-007"]


def test_prepare_treats_미착수_보류_지연_as_incomplete(tmp_path: Path):
    """Carry-forward rule: anything that isn't 완료 or 취소 is incomplete."""
    prep = _load_prepare(tmp_path)
    team = _team()
    storage_root = tmp_path / "store"

    _seed_prev_week_part(
        storage_root, team.name, "2026-W17",
        group_name="사업그룹", part_name="전략기획",
        rows_이번주=[
            {"업무ID": "W17-A", "상태": "미착수", "업무_지시내용": "a"},
            {"업무ID": "W17-B", "상태": "보류", "업무_지시내용": "b"},
            {"업무ID": "W17-C", "상태": "지연", "업무_지시내용": "c"},
        ],
    )

    prep.prepare_new_week(
        storage_root=storage_root, team=team,
        current_week="2026-W18", prev_week="2026-W17",
    )
    wb_new = load_workbook(storage_root / team.name / "2026-W18" / "전략기획.xlsx")
    carried_ids = sorted(
        r[0] for r in wb_new["지난주"].iter_rows(min_row=2, values_only=True) if r[0]
    )
    assert carried_ids == ["W17-A", "W17-B", "W17-C"]


def test_prepare_carry_forward_includes_prev_week_지난주_rows(tmp_path: Path):
    """Items already in prev week's 지난주 (carried from W16 to W17) and STILL
    incomplete must keep being carried — they don't fall off after one week."""
    prep = _load_prepare(tmp_path)
    team = _team()
    storage_root = tmp_path / "store"

    _seed_prev_week_part(
        storage_root, team.name, "2026-W17",
        group_name="사업그룹", part_name="전략기획",
        rows_이번주=[
            {"업무ID": "W17-N", "상태": "진행중", "업무_지시내용": "new this week"},
        ],
        rows_지난주=[
            {"업무ID": "W16-X", "상태": "진행중", "업무_지시내용": "old, still going"},
            {"업무ID": "W16-Y", "상태": "완료", "업무_지시내용": "old, finally done"},
        ],
    )

    prep.prepare_new_week(
        storage_root=storage_root, team=team,
        current_week="2026-W18", prev_week="2026-W17",
    )
    wb_new = load_workbook(storage_root / team.name / "2026-W18" / "전략기획.xlsx")
    carried_ids = sorted(
        r[0] for r in wb_new["지난주"].iter_rows(min_row=2, values_only=True) if r[0]
    )
    assert carried_ids == ["W16-X", "W17-N"], (
        f"W16-Y was 완료 → archive only, W16-X + W17-N → carry. Got {carried_ids}"
    )


def test_prepare_idempotent_on_existing_week(tmp_path: Path):
    """Running prepare twice for the same week must not double-add rows."""
    prep = _load_prepare(tmp_path)
    team = _team()
    storage_root = tmp_path / "store"

    _seed_prev_week_part(
        storage_root, team.name, "2026-W17",
        group_name="사업그룹", part_name="전략기획",
        rows_이번주=[
            {"업무ID": "W17-1", "상태": "진행중", "업무_지시내용": "ongoing"},
        ],
    )
    prep.prepare_new_week(
        storage_root=storage_root, team=team,
        current_week="2026-W18", prev_week="2026-W17",
    )
    prep.prepare_new_week(
        storage_root=storage_root, team=team,
        current_week="2026-W18", prev_week="2026-W17",
    )
    wb = load_workbook(storage_root / team.name / "2026-W18" / "전략기획.xlsx")
    carried_ids = [r[0] for r in wb["지난주"].iter_rows(min_row=2, values_only=True) if r[0]]
    assert carried_ids == ["W17-1"], (
        f"second prepare() should be a no-op for already-prepared rows; got {carried_ids}"
    )
