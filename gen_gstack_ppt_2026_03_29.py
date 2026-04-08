from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import json
from pathlib import Path

BASE = Path(__file__).resolve().parent
DATA_PATH = BASE / 'gstack_slides_2026_03_29.json'
OUTPUT_PATH = BASE / 'gstack-explained-2026-03-29.pptx'

with open(DATA_PATH, 'r', encoding='utf-8') as f:
    data = json.load(f)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(248, 250, 252)
NAVY = RGBColor(15, 23, 42)
BLUE = RGBColor(37, 99, 235)
SLATE = RGBColor(71, 85, 105)
LIGHT = RGBColor(226, 232, 240)
WHITE = RGBColor(255, 255, 255)


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header_band(slide, title_text):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.8))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    tf = band.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title_text
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = WHITE


def add_footer(slide, text='Gstack 설명 자료 | 2026-03-29'):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(6.85), Inches(12.2), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(10)
    p.font.color.rgb = SLATE


def add_bullets_slide(title_text, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header_band(slide, title_text)

    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.15), Inches(12.1), Inches(5.35))
    panel.fill.solid()
    panel.fill.fore_color.rgb = WHITE
    panel.line.color.rgb = LIGHT

    body = slide.shapes.add_textbox(Inches(0.95), Inches(1.5), Inches(11.3), Inches(4.8))
    tf = body.text_frame
    tf.word_wrap = True
    tf.clear()

    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(24)
        p.font.color.rgb = NAVY
        p.space_after = Pt(12)
        p.bullet = True

    add_footer(slide)


# Title slide
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, RGBColor(239, 246, 255))
hero = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(0.8), Inches(11.9), Inches(5.5))
hero.fill.solid()
hero.fill.fore_color.rgb = WHITE
hero.line.color.rgb = LIGHT

accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.8), Inches(11.9), Inches(0.22))
accent.fill.solid()
accent.fill.fore_color.rgb = BLUE
accent.line.fill.background()

box = slide.shapes.add_textbox(Inches(1.0), Inches(1.35), Inches(10.8), Inches(1.4))
tf = box.text_frame
p = tf.paragraphs[0]
p.text = data['title']
p.font.size = Pt(28)
p.font.bold = True
p.font.color.rgb = NAVY

p2 = tf.add_paragraph()
p2.text = data['subtitle']
p2.font.size = Pt(18)
p2.font.color.rgb = SLATE
p2.space_before = Pt(12)

p3 = tf.add_paragraph()
p3.text = '10장 요약 자료'
p3.font.size = Pt(16)
p3.font.color.rgb = BLUE
p3.space_before = Pt(18)

source_box = slide.shapes.add_textbox(Inches(1.0), Inches(4.75), Inches(10.8), Inches(1.1))
stf = source_box.text_frame
sp = stf.paragraphs[0]
sp.text = '주요 참고: GitHub 공식 저장소와 README 기반 요약'
sp.font.size = Pt(14)
sp.font.color.rgb = SLATE

add_footer(slide, 'Gstack 설명 자료 | Garry Tan / GitHub 공식 자료 기반')

for slide_data in data['slides']:
    add_bullets_slide(slide_data['title'], slide_data['bullets'])

prs.save(str(OUTPUT_PATH))
print(str(OUTPUT_PATH))
